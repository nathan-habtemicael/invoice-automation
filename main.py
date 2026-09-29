import imaplib
import email
import json
import logging
import pdfplumber
import os
import re
import openpyxl
from openpyxl.styles import Font
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Border, Side
from groq import Groq
from dotenv import load_dotenv

        # Logging mit Zeitstempel
logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")

        # Zugangsdaten aus .env laden
load_dotenv()

        # Groq-Client initialisieren
GROQ_API_KEY = os.getenv("GROQ_API_KEY")
if not GROQ_API_KEY:
    raise RuntimeError("GROQ_API_KEY muss in der .env gesetzt sein")
GROQ_MODEL = os.getenv("GROQ_MODEL", "openai/gpt-oss-120b")
groq_client = Groq(api_key=GROQ_API_KEY)

ERLAUBTE_ENDUNGEN = (".pdf", ".txt", ".html")


        # UID-Handling
def lade_uid_max(dateiname="uid_max.txt"):
    try:
        with open(dateiname, "r") as f:
            return int(f.read().strip())        #strip() entfernt lehrzeichenartige Zeichen vor und nach dem String
    except (FileNotFoundError, ValueError):
        return 0


def speichere_uid_max(uid, dateiname="uid_max.txt"):
    with open(dateiname, "w") as f:
        f.write(str(uid))


        # E-Mail-Suche vorbereiten
def search_email(uid_max, criteria):
    search_strings = []
    for criteria_dict in criteria:
        search_parts = []
        for key, value in criteria_dict.items():
            search_parts.append(f'{key} "{value}"')
        search_parts.append(f'UID {uid_max + 1}:*')
        search_strings.append(f'({" ".join(search_parts)})')
    return search_strings


        # Excel-Datei erstellen
def create_excel_table(alle_rechnungen, excel_file="Invoices.xlsx"):
    if os.path.exists(excel_file):
        # Bestehende Datei laden
        wb = openpyxl.load_workbook(excel_file)
        ws = wb.active
        start_row = ws.max_row + 1  # neue Daten darunter schreiben
    else:
        # Neue Datei erstellen
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Invoices"

        # Überschriften schreiben
        ws.cell(row=1, column=1, value="Datum")
        ws.cell(row=1, column=2, value="Firma")
        ws.cell(row=1, column=3, value="Rechnungsnummer")
        ws.cell(row=1, column=4, value="Zwischensumme")
        ws.cell(row=1, column=5, value="Gesamtsumme")

        # Spaltenbreite definieren
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 17
        ws.column_dimensions['E'].width = 17

        # Überschriften formatieren (fett + hellgrün)
        fill_color = PatternFill(fill_type="solid", fgColor="8EE53F")
        for col in range(1, 6):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True, size=12)
            cell.fill = fill_color

        start_row = 2  # ab zweiter Zeile Daten einfügen

        # Neue Daten anhängen
    for row_num, rechnung in enumerate(alle_rechnungen, start=start_row):
        ws.cell(row=row_num, column=1,
                value=rechnung.get("Datum").strftime("%d.%m.%Y") if rechnung.get("Datum") else "")
        ws.cell(row=row_num, column=2, value=rechnung.get("Firma", ""))
        ws.cell(row=row_num, column=3, value=rechnung.get("Rechnungsnummer", ""))

        zelle_zwischensumme = ws.cell(row=row_num, column=4, value=rechnung.get("Zwischensumme", 0.0))
        zelle_zwischensumme.number_format = '#,##0.00 "€"'

        zelle_gesamtsumme = ws.cell(row=row_num, column=5, value=rechnung.get("Gesamtbetrag", 0.0))
        zelle_gesamtsumme.number_format = '#,##0.00 "€"'


    dünne_linien = Side(border_style="thin", color="000000")
    max_row = ws.max_row
    max_col = 5

    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = Border(
                left=dünne_linien,
                right=dünne_linien,
                top=dünne_linien,
                bottom=dünne_linien
            )

    wb.save(excel_file)
    logging.info(f"Excel-Datei '{excel_file}' wurde aktualisiert.")

        # Daten extrahieren
SYSTEM_PROMPT = """Du extrahierst Rechnungsdaten aus Text. Antworte ausschließlich mit einem JSON-Objekt
mit genau diesen Feldern, ohne Erklärungen und ohne Markdown:
{
  "datum": "TT.MM.JJJJ" oder null,
  "rechnungsnummer": string oder null,
  "firma": string oder null,
  "zwischensumme": Zahl oder null,
  "gesamtsumme": Zahl oder null
}
Beträge als Zahl mit Punkt als Dezimaltrennzeichen, ohne Währungssymbol.
Erfinde keine Werte: Steht ein Wert nicht im Text, setze null."""


def parse_datum(wert):
    if not wert:
        return None
    try:
        return datetime.strptime(str(wert).strip(), "%d.%m.%Y")
    except ValueError:
        logging.warning(f"Ungültiges Datum von der KI: {wert!r}")
        return None


def parse_betrag(wert):
    if wert is None:
        return 0.0
    if isinstance(wert, (int, float)) and not isinstance(wert, bool):
        return float(wert)
    try:
        text = re.sub(r"[^\d,.\-]", "", str(wert))
        if "," in text and "." in text:
            # Das zuletzt vorkommende Zeichen ist das Dezimaltrennzeichen
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        else:
            text = text.replace(",", ".")
        return float(text)
    except ValueError:
        logging.warning(f"Ungültiger Betrag von der KI: {wert!r}")
        return 0.0


def extract_invoice_data_ki(text):
    try:
        antwort = groq_client.chat.completions.create(
            model=GROQ_MODEL,
            temperature=0.1,
            max_tokens=1000,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": f"Rechnungstext:\n{text[:6000]}"},
            ],
        )
        inhalt = antwort.choices[0].message.content
        logging.info(f"Antwort der KI: {inhalt}")
        daten = json.loads(inhalt)
    except json.JSONDecodeError as e:
        logging.error(f"KI-Antwort ist kein gültiges JSON: {e}")
        return None
    except Exception as e:
        logging.error(f"Fehler bei der Groq-Anfrage: {e}")
        return None

    if not isinstance(daten, dict):
        logging.error(f"KI-Antwort ist kein JSON-Objekt: {daten!r}")
        return None

    return {
        "Datum": parse_datum(daten.get("datum")),
        "Rechnungsnummer": str(daten.get("rechnungsnummer") or ""),
        "Firma": str(daten.get("firma") or ""),
        "Zwischensumme": parse_betrag(daten.get("zwischensumme")),
        "Gesamtbetrag": parse_betrag(daten.get("gesamtsumme")),
    }


        # Dateinamen aus E-Mails bereinigen (Schutz vor Path-Traversal)
def sichere_dateiname(name):
    name = os.path.basename(name.replace("\\", "/"))
    name = re.sub(r"[^A-Za-z0-9ÄÖÜäöüß._\- ]", "_", name).strip()
    if name in ("", ".", ".."):
        return None
    return name


        # Anhang verarbeiten
def process_attachment(part, save_dir="anhänge"):
    original = part.get_filename()
    dateiname = sichere_dateiname(original) if original else None
    if not dateiname or not dateiname.lower().endswith(ERLAUBTE_ENDUNGEN):
        logging.info(f"Anhang ignoriert (nicht erlaubter Dateiname/Typ): {original!r}")
        return None

    os.makedirs(save_dir, exist_ok=True)
    filepath = os.path.join(save_dir, dateiname)

    with open(filepath, 'wb') as f:
        f.write(part.get_payload(decode=True))
    logging.info(f"Anhang {dateiname} gespeichert")

    try:
        text = ''
        endung = dateiname.lower()
        if endung.endswith('.pdf'):
            with pdfplumber.open(filepath) as pdf:
                text = ''.join(page.extract_text() or '' for page in pdf.pages)
        elif endung.endswith('.txt'):
            with open(filepath, 'r', encoding='utf-8') as f:
                text = f.read()
        elif endung.endswith('.html'):
            with open(filepath, 'r', encoding='utf-8') as f:
                text = re.sub(r'<[^>]+>', '', f.read()).strip()

        rechnung = extract_invoice_data_ki(text)
        if rechnung:
            alle_rechnungen.append(rechnung)
            logging.info(f"Rechnungsdaten extrahiert: {rechnung}")
        return rechnung

    except Exception as e:
        logging.error(f"Fehler beim Verarbeiten von {dateiname}: {e}")
        return None


        # Hauptprogramm
server = "imap.gmail.com"
email_user = os.getenv("EMAIL_USER")
email_pass = os.getenv("EMAIL_PASS")
sender_email = os.getenv("SENDER_EMAIL")
if not email_user or not email_pass or not sender_email:
    raise RuntimeError("EMAIL_USER, EMAIL_PASS und SENDER_EMAIL müssen in der .env gesetzt sein")
email_format = 'RFC822'
criteria = [{"FROM": sender_email}, {"FROM": "service@paypal.de"}]
save_dir = "anhänge"
alle_rechnungen = []
uid_max = lade_uid_max()
uid_max_start = uid_max

mail = imaplib.IMAP4_SSL(server)
mail.login(email_user, email_pass)
mail.select("INBOX")
logging.info("IMAP-Server verbunden")

search_str = search_email(uid_max, criteria)

for search in search_str:
    result, data = mail.uid('search', None, search)
    uid_list = data[0].decode().split() if data[0] else []

    if not uid_list:
        logging.info(f"Keine neuen E-Mails für {search}")
        continue

    for single_uid in uid_list:
        uid_int = int(single_uid)
        # IMAP liefert bei "UID n:*" immer mindestens die höchste UID, auch wenn sie < n ist
        if uid_int <= uid_max_start:
            continue
        if uid_int > uid_max:
            uid_max = uid_int  # Max UID aktualisieren

        result, msg_data = mail.uid('fetch', single_uid, email_format)
        raw_email = msg_data[0][1]
        email_message = email.message_from_bytes(raw_email)

        for part in email_message.walk():
            if part.get_content_disposition() == 'attachment':
                process_attachment(part, save_dir)

# Excel speichern
if alle_rechnungen:
    create_excel_table(alle_rechnungen)
else:
    logging.info("Keine Rechnungen gefunden")

# UID speichern und Verbindung schließen
speichere_uid_max(uid_max)
mail.logout()
logging.info("Verbindung geschlossen")
